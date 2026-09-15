"""读取车辆数据源，生成独立 JSON；不修改 Excel/CSV，不控制游戏。"""
import argparse
import csv
import hashlib
import json
import re
import unicodedata
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
LEAGUES = ["青铜", "白银", "黄金", "白金", "翡翠", "钻石", "精英", "宗师", "传奇"]
ORDER_SHEET = "！复制到脚本选车页"

def key(title):
    text = unicodedata.normalize("NFKD", str(title).casefold())
    return re.sub(r"[^a-z0-9\u4e00-\u9fff]", "", text)

def vehicle_id(title):
    return "car_" + hashlib.sha256(key(title).encode()).hexdigest()[:16]

def sequence(value):
    tokens = [part.strip() for part in str(value or "").split(",") if part.strip()]
    if not tokens or not all(token.isdigit() and int(token) > 0 for token in tokens):
        raise ValueError(f"非法序列: {value!r}")
    values = [int(token) for token in tokens]
    if len(values) != len(set(values)):
        raise ValueError(f"重复位置: {value!r}")
    return values

def build(source_dir):
    with (source_dir / "国服_a9mmgj_top.csv").open(encoding="utf-8-sig", newline="") as stream:
        catalog = [{"id": vehicle_id(r["title"]), "title": r["title"], "class": r["class"], "league": r["league"]}
                   for r in csv.DictReader(stream)]
    by_name = {}
    for record in catalog:
        if record["league"] not in LEAGUES or record["class"] not in set("DCBASR"):
            raise ValueError(f"目录字段非法: {record}")
        name = key(record["title"])
        if name in by_name:
            raise ValueError(f"目录名称重复: {record['title']}")
        by_name[name] = record

    aliases_path = source_dir / "vehicle_name_aliases.json"
    aliases = json.loads(aliases_path.read_text(encoding="utf-8")) if aliases_path.exists() else {}
    aliases = {key(source): target for source, target in aliases.items()}
    for target in aliases.values():
        if key(target) not in by_name:
            raise ValueError(f"别名目标不在基础目录: {target}")

    workbook = openpyxl.load_workbook(source_dir / "多人选车.xlsx", data_only=True, read_only=True)
    issues, groups = [], []
    try:
        order_rows = list(workbook[ORDER_SHEET].iter_rows(values_only=True))
        if list(order_rows[0][:2]) != ["段位", "序号"]:
            raise ValueError("选车顺序页表头改变")
        league_rows = {r[0]: (i, r) for i, r in enumerate(order_rows[1:], 2) if r[0] in LEAGUES}
        if len(league_rows) != len(LEAGUES) or sum(r[0] in LEAGUES for r in order_rows[1:]) != len(LEAGUES):
            raise ValueError("选车顺序缺少段位")
        for league in LEAGUES:
            sheet_rows = list(workbook[league].iter_rows(values_only=True))
            if list(sheet_rows[0][:9]) != ["段位", "等级", "名称", "性能分", "昵称", "备注", "位置", "符号", "优先级"]:
                raise ValueError(f"{league} 表头改变")
            positions = {}
            for row_number, row in enumerate(sheet_rows[1:], 2):
                if not row[2]:
                    continue
                if row[0] != league or not isinstance(row[6], (int, float)) or int(row[6]) != row[6]:
                    raise ValueError(f"{league}!G{row_number}: 段位或位置非法")
                position = int(row[6])
                if position <= 0 or position in positions:
                    raise ValueError(f"{league} 位置重复或非法: {position}")
                positions[position] = (row_number, row)
            order_row_number, order = league_rows[league]
            values = sequence(order[2] if order[2] else order[1])
            if order[1] and sequence(order[1]) != values:
                raise ValueError(f"{league}: 公式缓存与仅文本序列不同，请在 Excel 中更新")
            entries = []
            for rank, position in enumerate(values, 1):
                if position not in positions:
                    raise ValueError(f"{league}: 序列位置 {position} 没有对应车型")
                row_number, row = positions[position]
                source_title = str(row[2]).strip()
                match = by_name.get(key(aliases.get(key(source_title), source_title)))
                title = match["title"] if match else source_title
                entry = {
                    "order": rank, "title": title, "source_title": source_title,
                    "catalog_id": match["id"] if match else None,
                    "class": match["class"] if match else row[1], "source_class": row[1],
                    "nickname": str(row[4]) if row[4] is not None else None,
                    "source_position": position, "source_priority": row[8],
                    "source_row": row_number, "source_league": league, "notes": row[5],
                }
                if match is None:
                    issues.append({"kind": "unmatched_title", "league": league, "title": source_title, "source_row": row_number})
                elif match["league"] != league or match["class"] != row[1]:
                    issues.append({"kind": "catalog_difference", "resolution": "use_catalog", "league": league, "title": title,
                                   "workbook_class": row[1], "catalog_class": match["class"], "catalog_league": match["league"]})
                entries.append(entry)
            groups.append({"league": league, "source_sequence_cell": f"{ORDER_SHEET}!C{order_row_number}",
                           "source_positions": values, "vehicles": entries})
    finally:
        workbook.close()
    # CSV 决定实际定位段位。保留本段位原顺序，迁入条目按来源顺序追加。
    destinations = {league: [] for league in LEAGUES}
    migrations = []
    moved = []
    catalog_by_id = {record["id"]: record for record in catalog}
    for group in groups:
        for entry in group["vehicles"]:
            entry["source_order"] = entry["order"]
            target = catalog_by_id[entry["catalog_id"]]["league"] if entry["catalog_id"] else group["league"]
            if target == group["league"]:
                destinations[target].append(entry)
            else:
                moved.append((target, entry))
                migrations.append({"title": entry["title"], "from": group["league"], "to": target})
    for target, entry in moved:
        destinations[target].append(entry)
    source_groups = [{"league": group["league"], "source_sequence_cell": group["source_sequence_cell"],
                      "source_positions": group["source_positions"]} for group in groups]
    for group in groups:
        group.pop("source_sequence_cell")
        group.pop("source_positions")
        group["vehicles"] = destinations[group["league"]]
        for rank, entry in enumerate(group["vehicles"], 1):
            entry["order"] = rank
    rotation = {"schema_version": 2, "source": "多人选车.xlsx", "mode": "经典系列赛",
                "ordering": "per_league_source_sequence", "position_usage": "reference_only",
                "league_authority": "catalog", "class_authority": "catalog", "migration_ordering": "append_to_destination",
                "source_groups": source_groups, "groups": groups}
    report = {"catalog_records": len(catalog), "rotation_entries": sum(len(g["vehicles"]) for g in groups),
              "counts_by_league": {g["league"]: len(g["vehicles"]) for g in groups},
              "league_migrations": migrations, "issues": issues}
    return {"vehicle_catalog.json": {"schema_version": 1, "source": "国服_a9mmgj_top.csv", "vehicles": catalog},
            "multiplayer_rotation.json": rotation, "vehicle_import_report.json": report}

def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-dir", type=Path, default=ROOT / "data/sources")
    parser.add_argument("--output-dir", type=Path, default=ROOT / "data/generated")
    args = parser.parse_args()
    outputs = build(args.source_dir)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    for name, value in outputs.items():
        text = json.dumps(value, ensure_ascii=False, indent=2) + "\n"
        assert json.loads(text) == value
        (args.output_dir / name).write_text(text, encoding="utf-8")
    report = outputs["vehicle_import_report.json"]
    review = "# 车辆数据差异处理记录\n\n定位段位与车辆等级均以已审核 CSV 为准。下面保留 Excel 原值与 CSV 的差异供追溯；已匹配条目均采用 CSV 值。迁入车辆追加到目标段位队尾。\n\n"
    review += "| 车型 | Excel 段位 / 等级 | CSV 段位 / 等级 | 问题 |\n|---|---|---|---|\n"
    for issue in report["issues"]:
        if issue["kind"] == "unmatched_title":
            review += f"| {issue['title']} | {issue['league']} | 未匹配 | 名称需审核或基础目录缺车 |\n"
        else:
            review += f"| {issue['title']} | {issue['league']} / {issue['workbook_class']} | {issue['catalog_league']} / {issue['catalog_class']} | 已采用 CSV |\n"
    review += "\n已确认：One77 白金、Aglaia 宗师、Aeon E 传奇、5N 青铜、Panamera 黄金；5N 名称别名已匹配。\n"
    for migration in report["league_migrations"]:
        review += f"\n- {migration['title']}：{migration['from']} → {migration['to']}（已迁移）。\n"
    (args.output_dir / "vehicle_import_review.md").write_text(review, encoding="utf-8")
    print(json.dumps(outputs["vehicle_import_report.json"], ensure_ascii=False, indent=2))

if __name__ == "__main__":
    main()
